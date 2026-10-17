"""Bounded document extraction; uploaded content is untrusted reference material."""
import io,json,zipfile
from pathlib import Path
from uuid import uuid4,UUID
from app.config import DATA
MAX_BYTES=10*1024*1024
EXTENSIONS={'.txt','.md','.csv','.json','.pdf','.docx'}
def extract(name,data):
 ext=Path(name).suffix.lower()
 if ext not in EXTENSIONS:raise ValueError('支持 TXT、Markdown、CSV、JSON、PDF、Word（DOCX）')
 if not data:raise ValueError('文件为空')
 if len(data)>MAX_BYTES:raise ValueError('单个文件不能超过10MB')
 if ext=='.pdf':
  from pypdf import PdfReader
  doc=PdfReader(io.BytesIO(data))
  if len(doc.pages)>100:raise ValueError('PDF最多100页，请拆分后上传')
  text='\n'.join((p.extract_text() or '')[:12000] for p in doc.pages)[:50000]
 elif ext=='.docx':
  with zipfile.ZipFile(io.BytesIO(data)) as z:
   if sum(f.file_size for f in z.infolist())>40*1024*1024:raise ValueError('文档解压后过大')
  from docx import Document
  doc=Document(io.BytesIO(data));parts=[p.text for p in doc.paragraphs]
  parts += [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
  text='\n'.join(parts)[:50000]
 else:
  for encoding in ('utf-8-sig','gb18030'):
   try:text=data.decode(encoding);break
   except UnicodeDecodeError:continue
  else:raise ValueError('文本编码无法识别，请保存为UTF-8')
  text=text[:50000]
 if not text.strip():raise ValueError('未提取到可读文字；扫描件请先转成可复制文字的PDF或Word')
 return text
def save_upload(name,data):
 text=extract(name,data);folder=DATA/'uploads';folder.mkdir(parents=True,exist_ok=True)
 token=uuid4().hex;name=Path(name.replace('\\','/')).name[:160]
 record={'id':token,'name':name,'size':len(data),'text':text,'truncated':len(text)>=50000}
 (folder/(token+Path(name).suffix.lower())).write_bytes(data)
 (folder/(token+'.json.meta')).write_text(json.dumps(record,ensure_ascii=False),encoding='utf8')
 return {k:v for k,v in record.items() if k!='text'}
def attach(p,ids):
 records=list(p.get('attachments',[]));known={a['id'] for a in records}
 for token in ids:
  try:UUID(token)
  except ValueError:raise ValueError('附件标识无效')
  if token in known:continue
  path=DATA/'uploads'/(token+'.json.meta')
  if not path.exists():raise ValueError('附件不存在，请重新上传')
  records.append(json.loads(path.read_text(encoding='utf8')));known.add(token)
 if len(records)>10:raise ValueError('每个项目最多10份资料')
 if sum(len(a['text']) for a in records)>100000:raise ValueError('资料文字过多，请精简或拆分项目')
 p['attachments']=records
