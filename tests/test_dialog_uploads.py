import unittest,tempfile,asyncio,json
from pathlib import Path
from unittest.mock import patch
import httpx
from app import store
from app.main import app,TASKS
from app.schemas import DirectorReply
from app.uploads import extract
class DialogTests(unittest.IsolatedAsyncioTestCase):
 async def asyncSetUp(self):
  self.tmp=tempfile.TemporaryDirectory();root=Path(self.tmp.name);self.patches=[patch.multiple(store,DB=root/'test.sqlite',DATA=root),patch('app.uploads.DATA',root)]
  for p in self.patches:p.start()
  store.initialize();self.client=httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://localhost:8020',headers={'Origin':'http://localhost:8020'})
 async def asyncTearDown(self):
  while TASKS:await asyncio.gather(*list(TASKS),return_exceptions=True)
  await self.client.aclose()
  for p in reversed(self.patches):p.stop()
  self.tmp.cleanup()
 async def test_questions_answer_and_upload_reach_model_without_production(self):
  upload=await self.client.post('/api/uploads?name=brand.txt',content='产品提供库存查询。'.encode());self.assertEqual(upload.status_code,200)
  questions=[{'id':'style','question':'喜欢哪种风格？','options':['简洁','活力']}]
  seen=[]
  async def provider(_self,p):
   seen.append(json.loads(json.dumps(p)))
   return {'reply':'请选择风格','ready':False,'plan':None,'questions':questions if len(seen)==1 else []}
  with patch('app.deepseek.DeepSeekProvider.converse',new=provider),patch('app.main.produce') as produce:
   response=await self.client.post('/api/projects',json={'mode':'text','source':'做宣传片','attachment_ids':[upload.json()['id']],'background':True});pid=response.json()['id']
   while TASKS:await asyncio.gather(*list(TASKS))
   p=store.get(pid);self.assertEqual(p['status'],'awaiting_input');self.assertEqual(len(p['questions']),1)
   response=await self.client.post('/api/projects/'+pid+'/conversations',json={'text':'采用推荐项','clarification_answers':{'style':'简洁'}});self.assertEqual(response.status_code,202)
   while TASKS:await asyncio.gather(*list(TASKS))
   self.assertTrue(seen[-1]['clarification_done']);self.assertIn('库存查询',seen[-1]['attachments'][0]['text']);produce.assert_not_called()
 async def test_reject_bad_attachment_and_oversize(self):
  response=await self.client.post('/api/uploads?name=run.exe',content=b'bad');self.assertEqual(response.status_code,400)
  response=await self.client.post('/api/uploads?name=big.txt',content=b'x'*(10*1024*1024+1));self.assertEqual(response.status_code,413)
  response=await self.client.post('/api/projects',json={'mode':'text','source':'test','attachment_ids':['../../secret']});self.assertEqual(response.status_code,400)
 def test_max_three_questions(self):
  q={'id':'q','question':'test','options':['a','b']}
  with self.assertRaises(ValueError):DirectorReply.model_validate({'reply':'test','questions':[q]*4})
 def test_document_text_and_tables(self):
  import io
  from docx import Document
  doc=Document();doc.add_paragraph('公司介绍');doc.add_table(rows=1,cols=1).cell(0,0).text='产品功能';data=io.BytesIO();doc.save(data)
  self.assertIn('产品功能',extract('brand.docx',data.getvalue()))
